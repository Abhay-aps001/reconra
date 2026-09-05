from __future__ import annotations

from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook  # type: ignore[import-untyped]

from .validation import ImportValidationError, ParsedTable, safe_cell


def inspect_xlsx(filename: str, content: bytes) -> ParsedTable:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_vba=False)
        sheets = [sheet for sheet in workbook.worksheets if sheet.max_row > 1]
        if len(sheets) != 1:
            raise ImportValidationError("AMBIGUOUS_WORKSHEET")
        rows = list(sheets[0].iter_rows(values_only=True))
    except ImportValidationError:
        raise
    except Exception as error:
        raise ImportValidationError("MALFORMED_WORKBOOK") from error
    headers = tuple(safe_cell(value).strip() for value in rows[0])
    if not all(headers) or len(set(headers)) != len(headers):
        raise ImportValidationError("MALFORMED_WORKBOOK")
    parsed = tuple(
        {header: _cell(value) for header, value in zip(headers, row, strict=True)}
        for row in rows[1:]
    )
    return ParsedTable(filename, "xlsx", headers, parsed)


def _cell(value: object) -> str:
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return safe_cell(value)
